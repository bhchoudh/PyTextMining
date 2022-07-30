import nltk
import datetime
import pandas as pd
import os, os.path
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from docx import Document
from openpyxl import load_workbook
from clsBhasTextMine import clsBhasTextMine

#nltk.download('punkt')                              #download NLTK packages
#nltk.download('stopwords')                          #download NLTK packages
#filepath="c:\BhaskarCode\FILES\TextMining\Docs"
filepath="d:\BhaskarCode\journal\cosine\Docs"
#df=pd.read_excel("c:\BhaskarCode\FILES\TFIDF.xlsx")
timestamp =str(datetime.datetime.now().day)+str(datetime.datetime.now().hour)+ str(datetime.datetime.now().minute)
txtpreproc = clsBhasTextMine()  #calling common preprocessing function from class

# Defining reference dataset to compare Cosine similarity
refdata = "10 + years experience, with 2+ years in Telecom industry Excellent communication & presentation skill & customer facing experience Containers – Dockers, Kubernetes & Helm Software Development Python CICD/ DevOps / Automation - Jenkins , Ansible, Puppet  etc Consulting, presales & E2E architecting experience Opensource forum upstream contribution Ability to lead multi location virtual teams in matrix set up Excellent communication skill & customer facing experience Deployment experience in OpenStack integration with Juniper Contrail SDN Installing, configuring and troubleshooting various OpenStack services using FUEL GUI and CLI Linux system administration experience Containers Dockers, Kubernetes & Helm Openstack Deployment & system integration experience"
refdata = txtpreproc.myTextPreProcessor(refdata)
refdata = ' '.join(refdata)

ranking=[]
for file in os.listdir(filepath):                       #loop through all files in the folder
    fullfilepath=os.path.join(filepath,file)
    if os.path.isfile(fullfilepath):
        #os.path.normpath(fullfilepath)
        #Reading DOCx file, creating string by appending paragraphs & vectorizing
        doc = Document(fullfilepath)
        testdata =''
        for para in doc.paragraphs:
            testdata= testdata+' ' +para.text
        testdata= txtpreproc.myTextPreProcessor(testdata)
        testdata = ' '.join(testdata)
        corpus = [refdata, testdata]

        print("*****************************************" )
        vectorizer = TfidfVectorizer()
        vectors = vectorizer.fit_transform(corpus)
        df = pd.DataFrame(vectors.todense(), columns= vectorizer.get_feature_names())
        cosine= cosine_similarity(vectors)
        #print(df)
        #print(cosine)
        cosine = cosine [0][1]
        ranking.append([file, cosine])
    else:
        continue


dfrank=pd.DataFrame(ranking,columns=['Name', 'Rating'])
dfrank.sort_values("Rating", axis=0,ascending=False,inplace=True,kind="quicksort")
print(dfrank)
excelpath = "d:\BhaskarCode\journal\cosine\RankCosine3007.xlsx"
book = load_workbook(excelpath)     #excel needs to exist else it will give error
dfwrite = pd.ExcelWriter(excelpath)
dfwrite.book = book
dfrank.to_excel(dfwrite, sheet_name=timestamp, index=True)
dfwrite.save()

quit()

#####  Dt 30/07/22 ######
#This text preprocessing being a common function is shifted to class clsBhasTextMine
#This function can be removed from this file, kept for past reference only
def myTextPreProcessor(text):
    stopWords = set(nltk.corpus.stopwords.words('english'))
    tokenizer = nltk.tokenize.RegexpTokenizer(r'\w+')  # Get tokenizer with builtin cleaning
    text = tokenizer.tokenize(text)
    for word in text:
        if word in stopWords:
            text.remove(word)  # remove stopwords
        elif not word.isalpha():
            text.remove(word)  # remove special characters
    return text
